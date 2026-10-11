from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from patients.models import Patient
from visits.models import Visit
from django.db.models import Count, Sum, Max, Min
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

@login_required
def reports_view(request):
    doctor = request.user
    today = timezone.now().date()
    
    # Analytics data
    total_consultations = Visit.objects.filter(patient__doctor=doctor).count()
    
    total_revenue = Visit.objects.filter(patient__doctor=doctor).aggregate(Sum('payments__amount'))['payments__amount__sum'] or 0
    avg_price = total_revenue / total_consultations if total_consultations > 0 else 0
    
    # Highest revenue day
    highest_day = Visit.objects.filter(patient__doctor=doctor) \
        .values('visit_date__date') \
        .annotate(daily_rev=Sum('payments__amount')) \
        .order_by('-daily_rev').first()
        
    highest_revenue_day = highest_day['visit_date__date'] if highest_day else None
    highest_revenue_amount = highest_day['daily_rev'] if highest_day else 0

    context = {
        'total_consultations': total_consultations,
        'avg_price': avg_price,
        'highest_revenue_day': highest_revenue_day,
        'highest_revenue_amount': highest_revenue_amount
    }
    return render(request, 'reports/index.html', context)

@login_required
def export_patients_excel(request):
    doctor = request.user
    patients = Patient.objects.filter(doctor=doctor).annotate(
        visit_count=Count('visits'),
        first_visit=Min('visits__visit_date'),
        last_visit=Max('visits__visit_date'),
        total_paid=Sum('visits__payments__amount')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patients"

    # Headers
    headers = ['Patient ID', 'Patient Name', 'Age', 'Gender', 'Phone', 'First Visit', 'Last Visit', 'Total Visits', 'Total Paid']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')

    # Data
    for row_num, patient in enumerate(patients, 2):
        ws.cell(row=row_num, column=1, value=patient.patient_id)
        ws.cell(row=row_num, column=2, value=patient.full_name)
        ws.cell(row=row_num, column=3, value=patient.age)
        ws.cell(row=row_num, column=4, value=patient.gender)
        ws.cell(row=row_num, column=5, value=patient.phone)
        ws.cell(row=row_num, column=6, value=patient.first_visit.strftime('%Y-%m-%d') if patient.first_visit else '')
        ws.cell(row=row_num, column=7, value=patient.last_visit.strftime('%Y-%m-%d') if patient.last_visit else '')
        ws.cell(row=row_num, column=8, value=patient.visit_count)
        
        paid_cell = ws.cell(row=row_num, column=9, value=patient.total_paid or 0)
        paid_cell.number_format = '₹#,##0.00'

    # Auto-width
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    date_str = datetime.now().strftime('%Y-%m-%d')
    response['Content-Disposition'] = f'attachment; filename=doctor_patient_report_{date_str}.xlsx'
    wb.save(response)
    return response

@login_required
def export_visits_excel(request):
    doctor = request.user
    visits = Visit.objects.filter(patient__doctor=doctor).select_related('patient').order_by('-visit_date')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visits Report"

    headers = ['Patient', 'Visit Date', 'Diagnosis', 'Consultation Price', 'Doctor Notes']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')

    for row_num, visit in enumerate(visits, 2):
        ws.cell(row=row_num, column=1, value=visit.patient.full_name)
        ws.cell(row=row_num, column=2, value=visit.visit_date.strftime('%Y-%m-%d %H:%M'))
        ws.cell(row=row_num, column=3, value=visit.diagnosis)
        
        fee_cell = ws.cell(row=row_num, column=4, value=visit.consultation_fee)
        fee_cell.number_format = '₹#,##0.00'
        
        ws.cell(row=row_num, column=5, value=visit.doctor_notes)

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    date_str = datetime.now().strftime('%Y-%m-%d')
    response['Content-Disposition'] = f'attachment; filename=doctor_visit_report_{date_str}.xlsx'
    wb.save(response)
    return response
